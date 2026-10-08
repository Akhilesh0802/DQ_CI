from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.dataset import DatasetVersion
from app.models.profiling import ColumnMetric, ProfilingResult, QualityIssue
from app.services.column_profiler import METRIC_DEFS

HEADER_FILL = PatternFill("solid", fgColor="4F2D7F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ERROR_FILL = PatternFill("solid", fgColor="FDE2E2")
WARNING_FILL = PatternFill("solid", fgColor="FDF1DC")

SHEET_TITLES = {"NUMERIC": "Numeric", "DATE": "Date", "TEXT": "Text", "BOOLEAN": "Boolean"}


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, min_width=12, max_width=40):
    for idx, column_cells in enumerate(ws.columns, start=1):
        longest = max((len(str(c.value)) for c in column_cells if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(idx)].width = max(min_width, min(longest + 2, max_width))


def build_excel_report(db: Session, version: DatasetVersion) -> bytes:
    dataset = version.dataset
    result = db.query(ProfilingResult).filter(ProfilingResult.dataset_version_id == version.id).first()
    columns = db.query(ColumnMetric).filter(ColumnMetric.dataset_version_id == version.id).all()
    issues = db.query(QualityIssue).filter(QualityIssue.dataset_version_id == version.id).all()

    wb = Workbook()

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    type_counts = {}
    for c in columns:
        type_counts[c.inferred_type] = type_counts.get(c.inferred_type, 0) + 1
    rows = [
        ("File name", dataset.file_name),
        ("System", dataset.system or "-"),
        ("Reporting period", dataset.reporting_period or "-"),
        ("Version", version.version_number),
        ("Uploaded at", version.uploaded_at.strftime("%d-%m-%Y %H:%M:%S") if version.uploaded_at else "-"),
        ("Status", version.status.value if version.status else "-"),
        ("Total records", result.row_count if result else "-"),
        ("Total columns", len(columns)),
        ("Numeric columns", type_counts.get("NUMERIC", 0)),
        ("Date columns", type_counts.get("DATE", 0)),
        ("Text columns", type_counts.get("TEXT", 0)),
        ("Boolean columns", type_counts.get("BOOLEAN", 0)),
        ("Duplicate rows", result.duplicate_count if result else "-"),
        ("Errors", sum(1 for i in issues if i.severity == "error")),
        ("Warnings", sum(1 for i in issues if i.severity == "warning")),
    ]
    ws.append(["Metric", "Value"])
    _style_header(ws)
    for r in rows:
        ws.append(list(r))
    _autosize(ws)

    # ---- Per-type sheets (reference jaisa: row = column, col = metric)
    for type_key, sheet_title in SHEET_TITLES.items():
        group = [c for c in columns if c.inferred_type == type_key]
        if not group:
            continue
        sheet = wb.create_sheet(sheet_title)
        defs = METRIC_DEFS[type_key]
        sheet.append(["Column"] + (["Detected format"] if type_key == "DATE" else []) + [label for _, label in defs])
        _style_header(sheet)
        for c in group:
            row = [c.column_name] + ([c.detected_format] if type_key == "DATE" else [])
            row += [c.metrics.get(key) for key, _ in defs]
            sheet.append(row)
        sheet.freeze_panes = "B2"
        _autosize(sheet)

    # ---- Issues
    ws_issues = wb.create_sheet("Issues")
    ws_issues.append(["Severity", "Column", "Message"])
    _style_header(ws_issues)
    for i in issues:
        ws_issues.append([i.severity.upper(), i.column_name or "Overall", i.message])
        fill = ERROR_FILL if i.severity == "error" else WARNING_FILL
        for cell in ws_issues[ws_issues.max_row]:
            cell.fill = fill
    if not issues:
        ws_issues.append(["-", "-", "Koi issue nahi mila"])
    _autosize(ws_issues, max_width=70)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
